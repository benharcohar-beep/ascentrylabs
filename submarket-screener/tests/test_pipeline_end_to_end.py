"""End to end test of everything downstream of the network.

Builds a SYNTHETIC bundle (the numbers are invented and exist only to exercise
the code paths; nothing here is real data and nothing here reaches an output
the user would read), then runs scoring, the workbook writer and the one-pager
writer over it. This is the test that proves the interview demo works offline.
"""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from screener import excel, onepager, score
from screener.config import load_weights
from screener.metrics import registry

SYNTHETIC_UNITS = [
    ("5502500", "Alpha Village",  "55025", 43.10, -89.30),
    ("5502501", "Bravo Town",     "55025", 43.20, -89.45),
    ("5502502", "Charlie City",   "55021", 43.35, -89.20),
    ("5502503", "Delta Village",  "55045", 42.80, -89.60),
    ("5502504", "Echo Township",  "55049", 43.05, -89.75),
]


def _cell(value, *, missing_reason=""):
    return {
        "value": value,
        "source": "SYNTHETIC TEST DATA, not a real source",
        "vintage": "SYNTHETIC",
        "url": "https://example.invalid/synthetic",
        "retrieved_at": "2026-01-01T00:00:00Z",
        "notes": "Invented for testing. Never appears in a real run.",
        "missing_reason": missing_reason or ("" if value is not None else "synthetic gap"),
    }


@pytest.fixture()
def bundle() -> dict:
    scored_specs, context_specs, owner = registry()
    specs = {**scored_specs, **context_specs}

    values: dict[str, dict] = {}
    for i, (geoid, name, county, lat, lon) in enumerate(SYNTHETIC_UNITS):
        row = {}
        for j, (key, spec) in enumerate(specs.items()):
            # Leave a couple of holes on purpose so the MISSING handling is
            # exercised rather than assumed.
            if (i + j) % 11 == 0:
                row[key] = _cell(None)
            elif spec.scored:
                row[key] = _cell(float(10 * (i + 1) + j))
            else:
                row[key] = _cell(f"context {i}-{j}")
        values[geoid] = row

    return {
        "market": {
            "key": "synthetic_test", "name": "Synthetic Test Market",
            "short_name": "Synthetic", "geo_type": "county_subdivision",
            "geo_type_reason": "test only", "trade_area_note": "test only",
            "counties": [{"fips": "55025", "name": "Test County", "state": "55"}],
            "employment_centers": [
                {"name": "Test CBD", "lat": 43.07, "lon": -89.38, "note": "test"}
            ],
            "min_population": 2500, "max_distance_miles": 35,
            "target_submarkets": 5, "notes": "test only",
        },
        "built_at": "2026-01-01T00:00:00Z",
        "universe_size": 40, "in_range_size": 12,
        "units": [
            {"geoid": g, "name": n, "geo_type": "county_subdivision",
             "state_fips": "55", "county_fips": c, "county_name": "Test County",
             "lat": lat, "lon": lon, "land_area_sqmi": 12.0, "zctas": [("53500", 1.0)]}
            for g, n, c, lat, lon in SYNTHETIC_UNITS
        ],
        "values": values,
        "specs": {
            k: {"key": s.key, "label": s.label, "pillar": s.pillar,
                "higher_is_better": s.higher_is_better, "unit": s.unit,
                "decimals": s.decimals, "scored": s.scored,
                "description": s.description, "source_module": owner.get(k, "")}
            for k, s in specs.items()
        },
        "provenance": {},
        "failures": {"synthetic_source": "this is a test message"},
    }


def test_scoring_ranks_every_unit(bundle):
    weights = load_weights()
    ranked = score.score_market(bundle, weights)
    assert len(ranked) == len(SYNTHETIC_UNITS)
    assert [r.rank for r in ranked] == [1, 2, 3, 4, 5]
    totals = [r.total for r in ranked]
    assert all(t is not None for t in totals)
    assert totals == sorted(totals, reverse=True)
    for r in ranked:
        assert 0.0 <= r.total <= 100.0
        assert 0.0 <= r.coverage <= 1.0


def test_percentrank_matches_excel_definition():
    values = [10.0, 20.0, 20.0, 40.0]
    # Excel's PERCENTRANK.INC gives (count strictly below) / (n - 1).
    # Excel truncates to three significant digits by default, so 1/3 comes back
    # as 0.333 and not 0.333333. This function matches that, deliberately.
    assert score.percentrank_inc(values, 10.0) == pytest.approx(0.0)
    assert score.percentrank_inc(values, 20.0) == pytest.approx(33.3)
    assert score.percentrank_inc(values, 40.0) == pytest.approx(100.0)
    assert score.percentrank_inc([5.0], 5.0) is None


def test_weight_override_changes_the_order(bundle):
    weights = load_weights()
    a = score.score_market(bundle, weights, pillar_weights={"demand": 100})
    b = score.score_market(bundle, weights, pillar_weights={"supply": 100})
    assert [r.name for r in a] != [r.name for r in b]


def test_missing_values_never_become_zero(bundle):
    weights = load_weights()
    ranked = score.score_market(bundle, weights)
    for r in ranked:
        for ps in r.pillars.values():
            for ms in ps.metrics.values():
                if ms.raw is None:
                    assert ms.sub_score is None
                    assert ms.missing_reason


def test_workbook_writes_four_tabs_with_live_formulas(bundle, tmp_path: Path):
    weights = load_weights()
    ranked = score.score_market(bundle, weights)
    out = excel.build_workbook(bundle, weights, score.explain(bundle, ranked),
                               tmp_path / "test.xlsx")
    assert out.exists()

    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert wb.sheetnames == ["Ranking", "Raw data", "Scoring", "Sources and gaps"]

    scoring = wb["Scoring"]
    formulas = [c.value for row in scoring.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("PERCENTRANK.INC" in f for f in formulas)
    assert any("SUMPRODUCT" in f for f in formulas)

    ranking = wb["Ranking"]
    rank_formulas = [c.value for row in ranking.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("LARGE(" in f and "MATCH(" in f for f in rank_formulas)

    raw = wb["Raw data"]
    assert raw.cell(row=2, column=1).value == "Source"
    assert raw.cell(row=3, column=1).value == "Vintage"
    assert any(
        c.value == "MISSING"
        for row in raw.iter_rows(min_row=excel.RAW_DATA_START) for c in row
    )

    # The file must be a valid zip container, which is what makes it openable.
    with zipfile.ZipFile(out) as z:
        assert "xl/workbook.xml" in z.namelist()


def test_one_pager_is_self_contained(bundle, tmp_path: Path):
    weights = load_weights()
    ranked = score.score_market(bundle, weights)
    out = onepager.render(bundle, ranked, score.explain(bundle, ranked), weights,
                          tmp_path / "one.html")
    text = out.read_text()
    assert "<svg" in text
    assert "Synthetic Test Market" in text
    # No external asset may be referenced, or the page breaks offline.
    for bad in ("http://", "https://cdn", "<script", "<link"):
        assert bad not in text.replace("https://example.invalid", "")
    assert "MISSING" in text


def test_explanation_is_plain_sentences(bundle):
    weights = load_weights()
    ranked = score.score_market(bundle, weights)
    lines = score.explain(bundle, ranked)
    assert 2 <= len(lines) <= 5
    for line in lines:
        assert line.endswith(".")
        assert chr(0x2014) not in line and chr(0x2013) not in line


def test_bundle_round_trips_through_json(bundle, tmp_path: Path):
    path = tmp_path / "raw.json"
    path.write_text(json.dumps(bundle, default=str))
    again = json.loads(path.read_text())
    assert again["units"] == json.loads(json.dumps(bundle["units"], default=str))


# ---------------------------------------------------------------------------
# The workbook is only useful if its live formulas give the same answer as the
# Python scorer. If they drift, the interview demo contradicts itself. This
# test recalculates the saved .xlsx with an independent Excel formula engine
# and compares. It is skipped when that engine is not installed, because it is
# a development dependency and not needed to run the tool.
# ---------------------------------------------------------------------------
formulas = pytest.importorskip("formulas", reason="pip install formulas to run the parity check")


def _recalculate(path: Path) -> dict:
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()
    return {k.upper(): v for k, v in solution.items()}


def _scalar(value):
    try:
        return value.value[0, 0]
    except (AttributeError, IndexError, TypeError):
        return value


@pytest.mark.parametrize("method", ["percentile", "zscore"])
def test_excel_formulas_reproduce_the_python_ranking(bundle, tmp_path: Path, method):
    weights = load_weights()
    weights.scoring_method = method
    ranked = score.score_market(bundle, weights, method=method)

    name = f"parity_{method}.xlsx"
    out = excel.build_workbook(bundle, weights, score.explain(bundle, ranked),
                               tmp_path / name)
    cells = _recalculate(out)

    def ranking_cell(ref: str):
        return _scalar(cells.get(f"'[{name.upper()}]RANKING'!{ref}"))

    for i, result in enumerate(ranked):
        row = 6 + i           # ranking table starts on row 6
        assert str(ranking_cell(f"B{row}")).strip() == result.name, (
            f"row {row}: workbook says {ranking_cell(f'B{row}')!r}, "
            f"Python says {result.name!r}"
        )
        assert float(ranking_cell(f"C{row}")) == pytest.approx(result.total, abs=0.05)


# ---------------------------------------------------------------------------
# The parity fixture above is deliberately well behaved: every column strictly
# increases across rows. Real data is not like that, and the three shapes below
# are the ones that actually broke the workbook. Each is checked against the
# Python scorer.
#
# Ties are NOT tested here, on purpose. The `formulas` engine returns the
# mid-rank for a tied value in PERCENTRANK.INC (0.166 for {10,10,30,40}), while
# Excel returns the lowest tied rank (0.0), which is what screener/score.py
# implements. Testing ties here would fail against the engine while the
# workbook was right. The tie behaviour of our own function is pinned by
# test_percentrank_matches_excel_definition instead.
# ---------------------------------------------------------------------------


def _blank_column(bundle: dict, key: str, keep_geoids: tuple[str, ...] = ()) -> None:
    for unit in bundle["units"]:
        if unit["geoid"] not in keep_geoids:
            bundle["values"][unit["geoid"]][key] = _cell(None)


def _first_scored_keys(bundle: dict, n: int) -> list[str]:
    return [k for k, s in bundle["specs"].items() if s["scored"]][:n]


@pytest.mark.parametrize("shape", ["single_value_column", "empty_column", "sparse_row"])
def test_workbook_survives_the_shapes_that_occur_in_real_data(bundle, tmp_path, shape):
    keys = _first_scored_keys(bundle, 3)

    if shape == "single_value_column":
        # ZORI goes MISSING wherever it covers under 25% of a submarket, so a
        # column with exactly one number in it is entirely plausible.
        # PERCENTRANK.INC over one value divides by zero.
        _blank_column(bundle, keys[0], keep_geoids=(SYNTHETIC_UNITS[0][0],))
    elif shape == "empty_column":
        # municipal_attitude before it has been researched.
        _blank_column(bundle, keys[1])
    else:
        # One submarket with almost nothing. It must not rank first.
        target = SYNTHETIC_UNITS[2][0]
        for key, spec in bundle["specs"].items():
            if spec["scored"] and key != keys[0]:
                bundle["values"][target][key] = _cell(None)
        bundle["values"][target][keys[0]] = _cell(10_000.0)

    weights = load_weights()
    ranked = score.score_market(bundle, weights)

    name = f"shape_{shape}.xlsx"
    out = excel.build_workbook(bundle, weights, score.explain(bundle, ranked),
                               tmp_path / name)
    cells = _recalculate(out)

    def ranking_cell(ref: str):
        return _scalar(cells.get(f"'[{name.upper()}]RANKING'!{ref}"))

    top_name = str(ranking_cell("B6")).strip()
    assert top_name, f"{shape}: the Ranking tab came back blank"

    for i, result in enumerate(ranked):
        assert str(ranking_cell(f"B{6 + i}")).strip() == result.name, (
            f"{shape}: row {6 + i} workbook={ranking_cell(f'B{6 + i}')!r} "
            f"python={result.name!r}"
        )

    if shape == "sparse_row":
        starved = next(r for r in ranked if r.geoid == SYNTHETIC_UNITS[2][0])
        assert starved.thin_data
        assert starved.rank == len(ranked), (
            "a submarket measured on one column out of many must not outrank "
            "submarkets that were measured properly"
        )
        assert top_name != starved.name


def test_a_thin_row_cannot_win_even_with_a_perfect_score(bundle):
    """The failure the cross-check found: one best-in-group figure scoring 100."""
    weights = load_weights()
    target = SYNTHETIC_UNITS[1][0]
    only = _first_scored_keys(bundle, 1)[0]
    for key, spec in bundle["specs"].items():
        if spec["scored"] and key != only:
            bundle["values"][target][key] = _cell(None)
    bundle["values"][target][only] = _cell(10_000_000.0)   # best in the group

    ranked = score.score_market(bundle, weights)
    winner = ranked[0]
    starved = next(r for r in ranked if r.geoid == target)

    assert starved.total == pytest.approx(100.0), "it still scores 100 on its one column"
    assert starved.thin_data
    assert winner.geoid != target
    assert starved.rank == len(ranked)


def test_coverage_ignores_pillars_that_did_not_back_the_score(bundle):
    weights = load_weights()
    weights.min_metrics_per_pillar = 2
    target = SYNTHETIC_UNITS[0][0]
    demand_keys = [k for k, s in bundle["specs"].items()
                   if s["scored"] and s["pillar"] == "demand"]
    for key in demand_keys[1:]:
        bundle["values"][target][key] = _cell(None)

    ranked = score.score_market(bundle, weights)
    row = next(r for r in ranked if r.geoid == target)
    assert row.pillars["demand"].score is None, "demand should have been dropped"
    # Coverage must not count the weight of a pillar that contributed nothing.
    other_demand_weight = sum(
        ps.weight_used for p, ps in row.pillars.items()
        if p == "demand" and ps.score is None
    )
    assert other_demand_weight > 0
    assert row.coverage < 1.0
    contributing = sum(
        ps.weight_used * ps.coverage for p, ps in row.pillars.items()
        if ps.score is not None
    )
    total_weight = sum(ps.weight_used for ps in row.pillars.values())
    assert row.coverage == pytest.approx(contributing / total_weight, abs=1e-6)


def test_min_metrics_per_pillar_applies_in_the_workbook_too(bundle, tmp_path):
    """The setting used to exist in Python only, so raising it desynced the two."""
    weights = load_weights()
    weights.min_metrics_per_pillar = 2

    target = SYNTHETIC_UNITS[0][0]
    demand_keys = [k for k, s in bundle["specs"].items()
                   if s["scored"] and s["pillar"] == "demand"]
    for key in demand_keys[1:]:
        bundle["values"][target][key] = _cell(None)

    ranked = score.score_market(bundle, weights)
    name = "min_metrics.xlsx"
    out = excel.build_workbook(bundle, weights, score.explain(bundle, ranked),
                               tmp_path / name)
    cells = _recalculate(out)

    for i, result in enumerate(ranked):
        row = 6 + i
        got = _scalar(cells.get(f"'[{name.upper()}]RANKING'!B{row}"))
        assert str(got).strip() == result.name
        total = _scalar(cells.get(f"'[{name.upper()}]RANKING'!C{row}"))
        assert float(total) == pytest.approx(result.total, abs=0.05), (
            f"{result.name}: workbook {total} vs python {result.total}"
        )
